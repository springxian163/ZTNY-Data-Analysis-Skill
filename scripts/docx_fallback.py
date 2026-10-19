# -*- coding: utf-8 -*-
"""
内置 HTML -> DOCX 转换器（仅依赖 python-docx）

为什么需要它：
    run.py 默认使用 WorkBuddy tencent-docx 插件的 html-to-docx 转换器，排版效果最好，
    但那个转换器只存在于装了 WorkBuddy 的机器上。为了让本技能在其他环境（服务器、
    同事电脑、CI）也能产出 Word，这里提供一套纯 python-docx 的降级实现。

覆盖范围（只针对本项目 html_writer 产出的 DOM 结构，不做通用转换器）：
    封面 section[role=cover] / 目录 div.doc-toc / 执行摘要 div.executive-summary
    （含 table.kpi-table） / 正文 main（h2 → Heading 1，h3 → Heading 2，
     table.data-table / ul.bullet-list / ol.num-list / p / p.doc-note）

排版约定：
    A4，页边距 2.54cm，正文宋体五号 1.5 倍行距，标题黑体，
    表格 Table Grid + 表头浅底，页脚居中页码。
"""
import os
import re
from html.parser import HTMLParser

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
GRAY = RGBColor(0x61, 0x61, 0x61)
DARK = RGBColor(0x21, 0x21, 0x21)
PRIMARY = RGBColor(0x15, 0x65, 0xC0)
HEADER_FILL = "F5F7FA"

VOID = {"br", "hr", "img", "meta", "link", "input", "col", "source", "area", "base"}


# --------------------------------------------------------------------------
# 极简 DOM
# --------------------------------------------------------------------------
class Node:
    __slots__ = ("tag", "attrs", "children", "parent")

    def __init__(self, tag, attrs=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.children = []
        self.parent = None

    def add(self, child):
        self.children.append(child)
        if isinstance(child, Node):
            child.parent = self

    def text(self):
        buf = []
        for c in self.children:
            buf.append(c if isinstance(c, str) else c.text())
        return "".join(buf)

    @property
    def cls(self):
        return (self.attrs.get("class") or "").strip()

    @property
    def role(self):
        return (self.attrs.get("role") or "").strip()

    def walk(self):
        for c in self.children:
            if isinstance(c, Node):
                yield c
                yield from c.walk()

    def find(self, tag=None, cls=None, role=None, stop_at_heading=False):
        """广度优先查找第一个匹配节点"""
        queue = [self]
        while queue:
            cur = queue.pop(0)
            for c in cur.children:
                if not isinstance(c, Node):
                    continue
                if tag and c.tag != tag:
                    ok = False
                else:
                    ok = True
                    if cls and cls not in c.cls.split():
                        ok = False
                    if role and c.role != role:
                        ok = False
                if ok:
                    return c
                queue.append(c)
        return None

    def find_all(self, tag=None, cls=None):
        out = []
        for n in self.walk():
            if tag and n.tag != tag:
                continue
            if cls and cls not in n.cls.split():
                continue
            out.append(n)
        return out


class _Builder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node("#root")
        self.stack = [self.root]

    def handle_starttag(self, tag, attrs):
        node = Node(tag, {k: (v or "") for k, v in attrs})
        self.stack[-1].add(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_startendtag(self, tag, attrs):
        self.stack[-1].add(Node(tag, {k: (v or "") for k, v in attrs}))

    def handle_endtag(self, tag):
        for i in range(len(self.stack) - 1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                return

    def handle_data(self, data):
        if data:
            self.stack[-1].add(data)


def parse_html(text):
    b = _Builder()
    b.feed(text)
    b.close()
    return b.root


# --------------------------------------------------------------------------
# 排版工具
# --------------------------------------------------------------------------
def _set_east_asia(run, font):
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), font)
    rfonts.set(qn("w:hAnsi"), font)
    rfonts.set(qn("w:eastAsia"), font)


def _shade(cell, hex_fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcpr.append(shd)


def _runs(node):
    """把节点展开成 [(text, bold, italic)]"""
    out = []

    def walk(n, bold=False, italic=False):
        for c in n.children:
            if isinstance(c, str):
                t = re.sub(r"\s+", " ", c.replace("\u00a0", " "))
                if t:
                    out.append((t, bold, italic))
                continue
            if c.tag in ("strong", "b"):
                walk(c, True, italic)
            elif c.tag in ("em", "i"):
                walk(c, bold, True)
            elif c.tag == "br":
                out.append(("\n", bold, italic))
            elif c.tag in ("script", "style"):
                continue
            else:
                walk(c, bold, italic)

    walk(node)
    # 合并首尾空白
    merged = []
    for t, b, i in out:
        if merged and merged[-1][1] == b and merged[-1][2] == i:
            merged[-1] = (merged[-1][0] + t, b, i)
        else:
            merged.append((t, b, i))
    if merged:
        merged[0] = (merged[0][0].lstrip(), merged[0][1], merged[0][2])
        merged[-1] = (merged[-1][0].rstrip(), merged[-1][1], merged[-1][2])
    return [m for m in merged if m[0]]


def _write_runs(par, runs, font=BODY_FONT, size=10.5, color=DARK):
    for text, bold, italic in runs:
        parts = text.split("\n")
        for i, seg in enumerate(parts):
            if i:
                par.add_run().add_break()
            if not seg:
                continue
            r = par.add_run(seg)
            r.bold = bold
            r.italic = italic
            r.font.size = Pt(size)
            r.font.color.rgb = color
            _set_east_asia(r, font)


def _heading(doc, text, level=1):
    sizes = {1: 16, 2: 14, 3: 12}
    before = {1: 18, 2: 12, 3: 8}
    after = {1: 10, 2: 8, 3: 4}
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before = Pt(before[level])
    pf.space_after = Pt(after[level])
    pf.keep_with_next = True
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(sizes[level])
    r.font.color.rgb = DARK if level > 1 else PRIMARY
    _set_east_asia(r, HEAD_FONT)
    return p


def _body_par(doc, node, size=10.5, first_indent=True):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)
    if first_indent:
        pf.first_line_indent = Pt(size * 2)
    _write_runs(p, _runs(node), size=size)
    return p


def _note_par(doc, node):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(8)
    pf.left_indent = Pt(12)
    _write_runs(p, _runs(node), size=9, color=GRAY)
    return p


def _table(doc, node, font_size=9):
    rows = node.find_all("tr")
    if not rows:
        return None
    ncol = max(len([c for c in r.children if getattr(c, "tag", "") in ("td", "th")])
               for r in rows)
    ncol = max(ncol, 1)
    table = doc.add_table(rows=0, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    thead = node.find("thead")
    head_rows = set(id(r) for r in thead.find_all("tr")) if thead else set()

    for rnode in rows:
        cells = [c for c in rnode.children if getattr(c, "tag", "") in ("td", "th")]
        if not cells:
            continue
        row = table.add_row()
        is_head = id(rnode) in head_rows or all(c.tag == "th" for c in cells)
        for i, cnode in enumerate(cells[:ncol]):
            cell = row.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.0
            align = (cnode.attrs.get("style") or "")
            if "right" in align:
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            elif "center" in align:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _write_runs(p, _runs(cnode), size=font_size,
                        color=DARK)
            if is_head:
                for r in p.runs:
                    r.bold = True
                _shade(cell, HEADER_FILL)
    return table


def _num_item(doc, node):
    """带编号的段落（编号已在文字内，只需做悬挂缩进）"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Pt(24)
    pf.first_line_indent = Pt(-12)
    pf.space_after = Pt(4)
    pf.line_spacing = 1.4
    _write_runs(p, _runs(node))
    return p


def _list(doc, node, ordered):
    items = [c for c in node.children if getattr(c, "tag", "") == "li"]
    for idx, li in enumerate(items, 1):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.left_indent = Pt(21)
        pf.first_line_indent = Pt(-10.5)
        pf.space_after = Pt(3)
        pf.line_spacing = 1.4
        prefix = f"{idx}. " if ordered else ""
        if prefix:
            r = p.add_run(prefix)
            r.font.size = Pt(10.5)
            _set_east_asia(r, BODY_FONT)
        _write_runs(p, _runs(li))


def _page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _footer_page_number(doc):
    footer = doc.sections[0].footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    run.font.size = Pt(9)
    run.font.color.rgb = GRAY


def _setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, side, Cm(2.54))

    st = doc.styles["Normal"]
    st.font.name = BODY_FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    st.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    st.paragraph_format.space_after = Pt(6)


# --------------------------------------------------------------------------
# 各区块
# --------------------------------------------------------------------------
def _segments(node):
    """按 <br> 切分成多行，用于封面主标题"""
    segs, cur = [], []

    def walk(n):
        for c in n.children:
            if isinstance(c, str):
                t = re.sub(r"\s+", " ", c.replace("\u00a0", " "))
                if t.strip():
                    cur.append(t)
            elif c.tag == "br":
                segs.append("".join(cur).strip())
                cur.clear()
            else:
                walk(c)

    walk(node)
    segs.append("".join(cur).strip())
    return [s for s in segs if s]


def _render_cover(doc, sec):
    for p in sec.children:
        if not isinstance(p, Node):
            continue
        cls = p.cls
        if "cover-category" in cls:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(80)
            r = par.add_run(p.text().strip())
            r.font.size = Pt(10.5)
            r.font.color.rgb = GRAY
            _set_east_asia(r, BODY_FONT)
            rpr = r._element.get_or_add_rPr()
            sp = OxmlElement("w:spacing")
            sp.set(qn("w:val"), "60")
            rpr.append(sp)
        elif p.tag == "h1" or "cover-title" in cls:
            for line in _segments(p):
                if not line.strip():
                    continue
                par = doc.add_paragraph()
                par.alignment = WD_ALIGN_PARAGRAPH.CENTER
                par.paragraph_format.space_before = Pt(18)
                par.paragraph_format.line_spacing = 1.4
                r = par.add_run(line.strip())
                r.bold = True
                r.font.size = Pt(26)
                r.font.color.rgb = DARK
                _set_east_asia(r, HEAD_FONT)
        elif "cover-subtitle" in cls:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(10)
            r = par.add_run(p.text().strip())
            r.font.size = Pt(13)
            r.font.color.rgb = PRIMARY
            _set_east_asia(r, BODY_FONT)
        elif "cover-meta" in cls:
            par = doc.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.space_before = Pt(6)
            r = par.add_run(p.text().strip())
            r.font.size = Pt(10)
            r.font.color.rgb = GRAY
            _set_east_asia(r, BODY_FONT)


def _render_toc(doc, toc):
    title = toc.find("p", cls="toc-title")
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(6)
    par.paragraph_format.space_after = Pt(12)
    r = par.add_run(title.text().strip() if title else "目　录")
    r.bold = True
    r.font.size = Pt(16)
    _set_east_asia(r, HEAD_FONT)

    for item in toc.find_all("p", cls="toc-item"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.3
        _write_runs(p, _runs(item), size=10.5)


def _render_main(doc, main):
    for node in main.children:
        if not isinstance(node, Node):
            txt = (node or "").strip()
            if txt:
                doc.add_paragraph(txt)
            continue
        tag = node.tag
        if tag == "h1":
            _heading(doc, node.text().strip(), 1)
        elif tag == "h2":
            _heading(doc, node.text().strip(), 1)
        elif tag == "h3":
            _heading(doc, node.text().strip(), 2)
        elif tag == "h4":
            _heading(doc, node.text().strip(), 3)
        elif tag == "p":
            if "num-item" in node.cls:
                _num_item(doc, node)
            elif "doc-note" in node.cls:
                _note_par(doc, node)
            else:
                _body_par(doc, node)
        elif tag == "table":
            _table(doc, node)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif tag == "ul":
            _list(doc, node, ordered=False)
        elif tag == "ol":
            _list(doc, node, ordered=True)
        elif tag == "blockquote":
            _note_par(doc, node)
        elif tag in ("div", "section"):
            _render_main(doc, node)


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------
def convert(html_path, docx_path, page_size="a4"):
    raw = open(html_path, encoding="utf-8").read()
    root = parse_html(raw)
    body = root.find("body") or root

    doc = Document()
    _setup(doc)

    cover = body.find("section", role="cover")
    if cover:
        _render_cover(doc, cover)
        _page_break(doc)

    toc = body.find("div", cls="doc-toc")
    if toc:
        _render_toc(doc, toc)
        _page_break(doc)

    summary = body.find("div", cls="executive-summary")
    if summary:
        h2 = summary.find("h2")
        if h2:
            _heading(doc, h2.text().strip(), 1)
        kpi = summary.find("table", cls="kpi-table")
        if kpi:
            _table(doc, kpi)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        body_div = summary.find("div", cls="summary-body")
        for p in (body_div.children if body_div else summary.children):
            if isinstance(p, Node) and p.tag == "p":
                _body_par(doc, p)

    main = body.find("main")
    if main:
        _render_main(doc, main)

    _footer_page_number(doc)

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    return docx_path


if __name__ == "__main__":
    import sys
    if len(sys.argv) < 3:
        raise SystemExit("用法: python docx_fallback.py in.html out.docx")
    print(convert(sys.argv[1], sys.argv[2]))
